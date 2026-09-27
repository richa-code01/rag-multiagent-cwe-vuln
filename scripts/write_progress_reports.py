"""Write ten short mentor progress reports as editable Word files.

Each report adds one step. It does not mention results from later reports.
Pictures are drawn with Pillow and embedded in the DOCX. Text stays real
Word text so the file can be edited.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "progress"
FIG = OUT / "figures"

NAVY = (27, 58, 75)
CREAM = (246, 241, 231)
RUST = (196, 92, 38)
INK = (28, 25, 21)
CARD = (255, 252, 247)
LINE = (217, 208, 195)
MUTED = (92, 86, 76)

TITLE = (
    "RAG-Augmented Multi-Agent LLM Framework for Explainable "
    "Software Vulnerability Detection Using CWE Knowledge Bases"
)
WHO = "Richa Verma (25MCSS02)    Advisor: Dr. Akshay Pandey"


def _font(size: int, bold: bool = False) -> ImageFont.ImageFont:
    names = [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/Library/Fonts/Arial Bold.ttf" if bold else "/Library/Fonts/Arial.ttf",
    ]
    for name in names:
        if Path(name).is_file():
            return ImageFont.truetype(name, size)
    return ImageFont.load_default()


def _canvas() -> tuple[Image.Image, ImageDraw.ImageDraw]:
    image = Image.new("RGB", (1600, 560), CREAM)
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 1600, 10), fill=NAVY)
    return image, draw


def _box(draw: ImageDraw.ImageDraw, xy: tuple[int, int, int, int], fill: tuple[int, int, int], title: str, body: str) -> None:
    draw.rounded_rectangle(xy, radius=18, fill=fill, outline=NAVY, width=3)
    x1, y1, x2, _y2 = xy
    draw.text((x1 + 24, y1 + 22), title, font=_font(28, bold=True), fill=NAVY)
    y = y1 + 72
    for line in body.split("\n"):
        draw.text((x1 + 24, y), line, font=_font(22), fill=INK)
        y += 32


def _arrow(draw: ImageDraw.ImageDraw, x1: int, x2: int, y: int) -> None:
    draw.line((x1, y, x2 - 18, y), fill=RUST, width=8)
    draw.polygon([(x2 - 18, y - 14), (x2, y), (x2 - 18, y + 14)], fill=RUST)


def _save(image: Image.Image, name: str) -> Path:
    FIG.mkdir(parents=True, exist_ok=True)
    path = FIG / name
    image.save(path, "PNG")
    return path


def fig_question() -> Path:
    image, draw = _canvas()
    _box(draw, (40, 80, 520, 480), CARD, "Java file", "A short method.\nMaybe a bug.\nMaybe safe code.")
    _arrow(draw, 540, 700, 280)
    _box(draw, (720, 70, 1560, 490), CARD, "A note a person can read", "Unsafe or safe\nWhich weakness (CWE)\nWhich lines\nWhy, and how to fix it")
    return _save(image, "01_question.png")


def fig_split() -> Path:
    image, draw = _canvas()
    _box(draw, (80, 70, 740, 490), CARD, "Train  ·  8 files", "4 unsafe\n4 safe\nWe use these to shape the tool.")
    _box(draw, (860, 70, 1520, 490), (255, 236, 224), "Test  ·  4 files", "2 unsafe\n2 safe\nHeld back for a first check.")
    return _save(image, "02_practice_set.png")


def fig_card() -> Path:
    image, draw = _canvas()
    _box(draw, (60, 50, 980, 510), CARD, "CWE-89  SQL injection", "Name and short description\nHow it happens\nHow to fix it\nParent and child weaknesses")
    _box(draw, (1040, 80, 1540, 230), (255, 236, 224), "Parent", "A broader group")
    _box(draw, (1040, 260, 1540, 470), CARD, "Child", "A more specific case")
    return _save(image, "03_knowledge.png")


def fig_search() -> Path:
    image, draw = _canvas()
    _box(draw, (40, 140, 460, 420), (255, 236, 224), "Question from code", "\"SQL string built\nby sticking text\ntogether\"")
    _arrow(draw, 480, 640, 280)
    _box(draw, (660, 40, 1560, 180), CARD, "1  CWE-89", "Best match")
    _box(draw, (660, 200, 1560, 340), CARD, "2  a related weakness", "Next match")
    _box(draw, (660, 360, 1560, 500), CARD, "3  another related page", "Still useful context")
    return _save(image, "04_retrieval.png")


def fig_answer() -> Path:
    image, draw = _canvas()
    rows = [
        "Decision     unsafe  /  safe  /  not sure",
        "CWE          id and name",
        "Lines        file and line numbers",
        "Why          one short cause",
        "Explain      plain words",
        "Fix          what to change",
    ]
    draw.rounded_rectangle((180, 40, 1420, 520), radius=18, fill=CARD, outline=NAVY, width=3)
    draw.text((220, 60), "Every answer uses the same fields", font=_font(30, bold=True), fill=NAVY)
    y = 130
    for row in rows:
        draw.text((240, y), row, font=_font(26), fill=INK)
        y += 58
    return _save(image, "05_answer_shape.png")


def fig_pipeline() -> Path:
    image, draw = _canvas()
    labels = [
        (40, "1  Mark lines", "Find risky spans.\nNot the final yes/no."),
        (420, "2  Search book", "Pick CWE pages\nthat fit the code."),
        (800, "3  Write note", "Fill the answer\nfrom a template."),
        (1180, "4  Check note", "Fields present.\nLines are real."),
    ]
    for x, title, body in labels:
        _box(draw, (x, 120, x + 360, 440), CARD, title, body)
    for x in (400, 780, 1160):
        _arrow(draw, x, x + 20, 280)
    return _save(image, "06_pipeline.png")


def fig_traps() -> Path:
    image, draw = _canvas()
    _box(draw, (60, 70, 740, 490), (255, 228, 224), "Simple checker", "Score 0 on 24 hard files.\nIt misses real bugs\nand also flags safe code.")
    _arrow(draw, 760, 900, 280)
    _box(draw, (920, 70, 1540, 490), (226, 240, 230), "Live model", "Score 0.923 on the same 24.\n2 false alarms.\n0 missed bugs.")
    return _save(image, "07_traps.png")


def fig_bars() -> Path:
    image, draw = _canvas()
    rows = [
        ("Juliet", 0.316, 20728),
        ("OWASP", 0.395, 2740),
        ("Securibench", 0.072, 119),
        ("Find Security Bugs", 0.435, 79),
        ("Vul4J", 0.244, 62),
        ("CVEfixes slice", 0.207, 92),
    ]
    draw.text((40, 24), "Simple checker score (F1) on six public Java sets", font=_font(28, bold=True), fill=NAVY)
    y = 80
    for name, score, n in rows:
        draw.text((40, y + 8), name, font=_font(22), fill=INK)
        width = int(700 * score / 0.5)
        draw.rounded_rectangle((420, y, 420 + max(width, 8), y + 42), radius=8, fill=RUST)
        draw.text((1140, y + 8), f"{score:.3f}    n={n}", font=_font(22), fill=INK)
        y += 72
    return _save(image, "08_six_suites.png")


def fig_pair() -> Path:
    image, draw = _canvas()
    _box(draw, (50, 80, 620, 420), (255, 228, 224), "Bad file", "Must be called unsafe.")
    _box(draw, (980, 80, 1550, 420), (226, 240, 230), "Good file", "Must be called safe.")
    draw.text((660, 200), "Both must\nbe right", font=_font(32, bold=True), fill=NAVY)
    draw.text((80, 450), "One right and one wrong  =  the pair is wrong", font=_font(24), fill=MUTED)
    return _save(image, "09_pairs.png")


def fig_console() -> Path:
    image, draw = _canvas()
    _box(draw, (40, 70, 700, 490), CARD, "Web page on this computer", "Pick a test set\nPick a setting\nSet a token budget\nPress start")
    _arrow(draw, 720, 880, 280)
    _box(draw, (900, 70, 1560, 490), CARD, "Same pipeline as the command line", "Mark lines\nSearch the CWE book\nWrite the note\nCheck the note")
    return _save(image, "10_console.png")


REPORTS: list[dict] = [
    {
        "slug": "01_question",
        "name": "What we want the tool to explain",
        "figure": fig_question,
        "intro": [
            "Many scanners only say that a piece of code is unsafe. The programmer still has to guess the kind of mistake and the fix.",
            "This note explains the problem we are working on. Later notes will describe the files, the knowledge book, the search, and the scores. Those pieces are not built in this note.",
        ],
        "did": [
            "We want a tool that reads a short Java method and writes a note a person can check.",
            "The note should say four things: whether the code is unsafe, which weakness it is, which lines matter, and how to fix it.",
            "The weakness name comes from CWE. CWE is a public list of common software mistakes. We use the official name so two people mean the same thing.",
            "We are studying six Java weaknesses only. They are listed in the table. They cover database queries, web pages, file paths, saved objects, passwords in code, and weak encryption.",
            "We are not trying to replace a commercial scanner. The aim is an explanation that can be traced back to a CWE page and to real lines of code.",
        ],
        "example": [
            "Suppose a method builds a database query by joining a username onto the text of the query. A yes-or-no scanner might only say \"unsafe\".",
            "The note we want is longer. It should say the code is unsafe, the weakness is CWE-89 SQL injection, the risky lines are the ones that build the query, and the fix is to use a prepared query so the username is data, not part of the command.",
            "If the same method uses a prepared query, the note should say the code is safe, and it should still name CWE-89 so a reader knows which risk was checked.",
        ],
        "caption": "The left box is the Java we read. The right box is the note we want back.",
        "headers": ["Weakness", "Plain meaning"],
        "rows": [
            ["CWE-89", "SQL injection. User text is stuck into a database query."],
            ["CWE-79", "Cross-site scripting. User text is shown as part of a web page."],
            ["CWE-22", "A file path can be steered outside the intended folder."],
            ["CWE-502", "Loading a saved object can run attacker-controlled data."],
            ["CWE-798", "A password is written directly in the code."],
            ["CWE-327", "The code uses weak encryption or a weak hash."],
        ],
        "table_note": "These six are the whole scope. A different kind of bug is outside this project, and we will not relabel it to force a match.",
        "outro": [
            "The next piece of work is a small set of Java examples, half unsafe and half safe, so the idea can be tried on real files.",
            "Until those files exist, there is nothing to score.",
        ],
        "not_claim": "This note has no test scores. It does not say the tool is better than any other scanner.",
    },
    {
        "slug": "02_practice_set",
        "name": "A small Java practice set and a first check",
        "figure": fig_split,
        "intro": [
            "The aim is a Java tool that explains a weakness in plain words, using a CWE name. This note is only about the first files and a very simple checker.",
            "Nothing here is a public benchmark. Juliet and OWASP are larger test sets used by other people. We have not used them yet.",
        ],
        "did": [
            "We wrote 12 short Java files, two for each of the six weaknesses: one unsafe and one safe. For example, one file builds a database query by sticking user text onto the query, and the paired file uses a safer call.",
            "Eight files are the practice split. Four files are held back. The held-back split has two unsafe files and two safe files. We do not tune the checker on those four.",
            "The checker looks for risky text patterns. It might see a query built with string joining, or a password written as a literal. It does not understand what the program does.",
            "On these 12 files it marked every unsafe file and did not mark a safe file. Precision means how many alarms were real bugs. Recall means how many real bugs we caught. Both are 1.000, so F1 is 1.000. F1 is one number that combines those two.",
            "A perfect score is easy to get when we wrote both the files and the patterns. It only shows that the checker matches this tiny set.",
        ],
        "example": [
            "One practice file builds a SQL string with the username inside it. The checker sees that pattern and marks the file unsafe. The paired file uses a prepared statement. The checker does not mark it.",
            "The four held-back files are the same idea for two other weaknesses: a password written in the source, and a weak hash. We did not change the patterns after seeing those four. They still came out right.",
            "That is why the table is balanced: every weakness in this set has one unsafe file and one safe file. The score cannot hide behind a set that is almost all unsafe.",
        ],
        "caption": "Eight files are for practice. Four files are held back for a first check.",
        "headers": ["Split", "Files", "Unsafe", "Safe"],
        "rows": [
            ["Train", "8", "4", "4"],
            ["Test", "4", "2", "2"],
            ["All", "12", "6", "6"],
        ],
        "table_note": "The score on all 12 files is precision 1.000, recall 1.000, F1 1.000. False alarms: 0. Missed bugs: 0.",
        "outro": [
            "The checker can say that a pattern fired. It still cannot explain the weakness in CWE words, because those words are not stored yet.",
            "The next piece of work is a small book of CWE facts: the name, what goes wrong, and a fix.",
        ],
        "not_claim": "This perfect score is not evidence about real projects. These 12 files are not Juliet and not the OWASP Benchmark.",
    },
    {
        "slug": "03_knowledge_book",
        "name": "A small book of CWE facts",
        "figure": fig_card,
        "intro": [
            "We already have 12 short Java files and a pattern checker that matches them. That checker can raise an alarm. It cannot yet say the official name of the weakness or how to fix it.",
            "This note describes the book we store those facts in. It does not change the detection score.",
        ],
        "did": [
            "CWE pages are published by MITRE. We do not copy the whole catalog. We keep the pages we need for the six weaknesses, plus some close neighbors, so a near miss can still be named.",
            "Each stored page has an id, a name, a short description, a longer description when MITRE provides one, links to parent and child weaknesses, and suggested fixes.",
            "The tool can look up one id, search by words such as \"cross site scripting\", and list neighbors. The pattern checker is not allowed to invent this text. If the answer names a CWE, that name has to come from this book.",
            "Take CWE-89 as the example in the picture. The page says SQL injection. It says what goes wrong when user text is joined into a query. It also points at related pages and at a safer style of query.",
            "A parent weakness is a broader group. A child weakness is a more specific case. We keep both so the explanation can stay honest when the code is close but not exact.",
        ],
        "example": [
            "Ask the book for CWE-89. It returns the name SQL injection, a short description, and a fix such as keeping user text out of the query command.",
            "Ask it for the neighbors of CWE-89. Parents are broader groups. Children are more specific kinds of injection. The answer can say \"this is SQL injection, which sits under this broader group\" without inventing a new name.",
            "Search the words \"cross site scripting\". The top hit should be the XSS page, CWE-79, not the SQL page. That is a lookup test, not a test of whether some Java file is unsafe.",
        ],
        "caption": "One stored page, CWE-89, and the kind of related pages we keep beside it.",
        "headers": ["On each page", "Why we keep it"],
        "rows": [
            ["Id and name", "So the answer uses the official name."],
            ["Description", "So the note can say what the weakness is."],
            ["Parents and children", "So a close match can still be explained."],
            ["Fixes", "So the answer can suggest a repair, not an attack."],
        ],
        "table_note": "This book is a subset. The full CWE list is much larger. We record that limit instead of pretending the book is complete.",
        "outro": [
            "The book can be queried by a person. The Java file does not search it by itself yet.",
            "The next piece of work is to turn the risky part of a Java file into a question and find the best CWE pages for that question.",
        ],
        "not_claim": "Storing the book does not make detection better. The score on the 12 practice files is unchanged.",
    },
    {
        "slug": "04_search",
        "name": "Finding the right CWE page",
        "figure": fig_search,
        "intro": [
            "We have short Java examples, a pattern checker, and a small book of CWE facts. This note is about search: given some Java, which book page should we show.",
            "Search is not the same as detection. A good search finds the right weakness page. It does not, by itself, decide that the code is unsafe.",
        ],
        "did": [
            "We wrote 18 labeled questions from the practice files. Each question has a known right CWE. We then ask three searches to rank the book.",
            "Word match (TF-IDF) looks for the same words in the question and on the page. It does not know that \"query built by joining strings\" means SQL injection unless those words overlap.",
            "Meaning match uses a small model called MiniLM. It can rank a page higher when the idea is similar, even if the words are not identical. If that model cannot be loaded, we fall back to word match and we write that down.",
            "The mix uses both rankings, plus two extra hints: the CWE id suggested by the pattern checker, and the parent or child pages of a hit. The mix is called hybrid search.",
            "R@1 means the right page was the first result. MiniLM was first on 17 of the 18 questions, which is 0.944. Word match was first on about 14 of 18, which is 0.778. On this small set the mix matched MiniLM. It did not beat it.",
        ],
        "example": [
            "The question is the idea \"a SQL string is built by sticking text together\". Word match may miss this if the CWE page says \"injection\" and the question says \"sticking text\". The meaning model can still rank CWE-89 first.",
            "If the pattern checker has already hinted CWE-89, the mix gives that hint a vote as well. Related pages, such as a child of CWE-89, can appear just below the best hit. They are context, not a second decision.",
            "If MiniLM cannot be downloaded, the same question is answered by word match only, and the result file says so. We do not quietly pretend the meaning model ran.",
        ],
        "caption": "A short question from the code is ranked against the book. The first hit is the page we would show.",
        "headers": ["Search", "Right page first (R@1)", "Questions"],
        "rows": [
            ["Word match (TF-IDF)", "0.778", "18"],
            ["Meaning match (MiniLM)", "0.944", "18"],
            ["Mix of both", "0.944", "18"],
        ],
        "table_note": "R@1 is not a detection score. It only says whether the correct CWE page was ranked first. These 18 questions come from our own files.",
        "outro": [
            "We can now retrieve a page. We still do not have one fixed shape for the written answer, so two runs could look different.",
            "The next piece of work is a single answer form: decision, CWE, lines, cause, explanation, and fix.",
        ],
        "not_claim": "This is not a public benchmark. A high R@1 here does not mean the tool finds bugs in Juliet or in real projects.",
    },
    {
        "slug": "05_answer_shape",
        "name": "One shape for every answer",
        "figure": fig_answer,
        "intro": [
            "We can mark risky lines, and we can search the CWE book. What was missing was a single shape for the written answer. Without that, two runs cannot be compared.",
            "This note describes that form. It does not add a new detection score.",
        ],
        "did": [
            "Every answer must fill the same fields. If a required field is missing, or if the text is not valid JSON, the answer is rejected.",
            "The decision is one of three words: unsafe, safe, or not sure. Not sure is allowed. We would rather see an honest abstention than a guessed yes.",
            "The CWE field must be an id and a name. The lines field must name the file and the line numbers, counted from 1, the same way an editor counts. The other fields are a short cause, a plain explanation, and a fix.",
            "There is no field for attack steps. The tool is not asked to write an exploit or a payload.",
            "We store two sample answers, one unsafe and one safe, and one broken sample. A small check accepts the good samples and rejects the broken one. That check does not need a live model.",
            "The pattern checker still only marks lines. The form is the contract for whoever writes the final note, whether that writer is a template or a model.",
        ],
        "example": [
            "A valid unsafe answer for the SQL example names the decision unsafe, the CWE as CWE-89, the lines that build the query, a cause such as \"user text is joined into the command\", a short explanation, and a fix that says to use a prepared query.",
            "A valid safe answer for the paired file uses the same fields. The decision is safe. It can still name CWE-89, because that is the risk that was checked.",
            "If the cited lines are blank, or the decision is a word we did not allow, the check fails. We keep that failure. We do not fill the hole by hand.",
        ],
        "caption": "Every run has to fill these fields. A missing field makes the answer invalid.",
        "headers": ["Field", "What it means"],
        "rows": [
            ["Decision", "Unsafe, safe, or not sure."],
            ["CWE", "Id and name from the book."],
            ["Lines", "File and line numbers, starting at 1."],
            ["Cause and explanation", "Why the lines are risky, in plain words."],
            ["Fix", "What a programmer should change."],
        ],
        "table_note": "The form is strict on purpose. Extra fields are not allowed, so the answer cannot hide an exploit in an extra box.",
        "outro": [
            "The pieces now exist separately: files, a checker, a book, search, and a form.",
            "The next piece of work is to run them in one order, still without calling a paid model.",
        ],
        "not_claim": "A valid form is not a correct decision. The checker can still be wrong about whether the code is unsafe.",
    },
    {
        "slug": "06_pipeline",
        "name": "The steps running in one order",
        "figure": fig_pipeline,
        "intro": [
            "The separate pieces are a pattern checker, a CWE book, search, and a fixed answer form. This note connects them into one run.",
            "No paid language model is used here. A template fills the form from the evidence. That template is a stand-in, so we can test the plumbing.",
        ],
        "did": [
            "The run always starts by marking risky lines. Those marks are evidence. They include the file, the line numbers, and a short reason. They are not the final yes or no.",
            "Next, search ranks CWE pages using the risky lines, not the whole file and not any hidden answer key. The pages include the description text, so a later writer can quote the book.",
            "A template then fills the answer form. If the evidence is empty, the template is more likely to say the code is safe. If evidence exists, it leans on that evidence. This is a rule, not a judgment.",
            "A checker then reads the form. The CWE must exist in the book. The cited lines must appear in the file. The fields must not contradict each other. A broken form fails.",
            "If the written decision disagrees with the pattern checker, we keep a warning. We do not force the decision to match the checker. We also save a confidence number. We do not use that number to skip any step.",
        ],
        "example": [
            "Take the SQL file again. Step 1 marks the lines that join the username into the query. Step 2 retrieves the CWE-89 page. Step 3, the template, fills the form from those marks and says the code is unsafe. Step 4 checks that CWE-89 is in the book and that the cited lines really are in the file.",
            "The paired safe file may produce no marks. The template then leans toward safe. The checker still requires a complete form.",
            "If the template said safe while step 1 had found marks, we would record a disagreement warning. We would not change the written decision to match the marks.",
        ],
        "caption": "Four steps, always in this order. The first step does not make the final decision.",
        "headers": ["Step", "What it is allowed to do"],
        "rows": [
            ["Mark lines", "Point at risky spans. Do not decide the case."],
            ["Search", "Pick CWE pages from the book."],
            ["Write", "Fill the form. Here, a template does this."],
            ["Check", "Reject a broken or unsupported form."],
        ],
        "table_note": "On the four held-back practice files, this template path still matches the labels. That only repeats the easy result from the tiny set. It is not a new public score.",
        "outro": [
            "The path is runnable from a command. The written note is still coming from rules, not from a model that can read a harder example.",
            "The next piece of work is a set of files written to fool the pattern checker, and a live model that has to explain them.",
        ],
        "not_claim": "The template is not the system we want to rely on. A saved confidence number is not a real probability.",
    },
    {
        "slug": "07_hard_traps",
        "name": "Hard examples and a live model",
        "figure": fig_traps,
        "intro": [
            "The pipeline can already mark lines, search the CWE book, fill the answer form, and check that form. On our first 12 files the simple checker looked perfect, because those files were easy.",
            "This note is about 24 harder files we wrote on purpose, and about a live language model. These files are still ours. They are not a public test set.",
        ],
        "did": [
            "Twelve of the new files are unsafe but quiet. Twelve are safe but look risky, often because a comment or a name mentions a flaw. A pattern checker that hunts for those words will get both kinds wrong.",
            "On the original text, the simple checker scores F1 = 0. It raises 12 false alarms and misses 12 bugs. F1 is the combined score of false alarms and missed bugs. Zero means it failed both ways.",
            "Before the model sees a file, we blank comments and rename obvious labels such as bad and good. That stops the question from containing the answer. After that cleaning, the template still scores F1 = 0. It has fewer false alarms, because comment-only hits disappear, but it still misses every bug.",
            "The live model is Groq, model openai/gpt-oss-20b. It writes the same answer form. On these 24 files its F1 is 0.923: 2 false alarms and no missed bugs.",
            "We ran the same 24 again with CWE search turned off. The F1 stayed 0.923. So, on this particular set, showing the model the book pages did not change detection. We do not treat that as proof that search is useless in general. We also do not treat it as proof that search helps.",
        ],
        "example": [
            "One safe file has a comment that says the code is flawed, but the code itself is fine. The simple checker reads the comment and raises a false alarm. After we blank comments, that alarm can disappear. The template still does not understand a quiet bug.",
            "One unsafe file never uses the obvious words. The pattern checker misses it. The live model reads the method and can still call it unsafe and name a CWE.",
            "Turning the book search off did not change the 0.923 score. The model was already getting the decision right from the code. The book may still matter for the explanation, but this test does not prove that.",
        ],
        "caption": "The simple checker fails these traps. The live model recovers most of them.",
        "headers": ["System", "F1", "False alarms", "Missed bugs", "Files"],
        "rows": [
            ["Simple checker, original text", "0.000", "12", "12", "24"],
            ["Template, after cleaning", "0.000", "6", "12", "24"],
            ["Live model", "0.923", "2", "0", "24"],
            ["Live model, search off", "0.923", "2", "0", "24"],
        ],
        "table_note": "The live model caught every unsafe file in this set and was wrong on 2 safe files. Precision is 0.857 and recall is 1.000. Those two numbers combine into F1 0.923.",
        "outro": [
            "This is a useful contrast: rules fail, the model mostly recovers. It is not a public result.",
            "The next measurement should be the simple checker on named public Java sets, reported separately from these 24 files.",
        ],
        "not_claim": "These 24 files are not Juliet. We do not say that CWE search improves detection.",
    },
    {
        "slug": "08_six_public_sets",
        "name": "Six public Java test sets",
        "figure": fig_bars,
        "intro": [
            "Our own 24 hard files showed that a live model can recover cases the pattern checker misses. Those files were written by us, so they are not enough.",
            "This note reports the same simple checker on six public Java sets that other people published. The live-model numbers on these sets are not in the table, and the reason is below.",
        ],
        "did": [
            "Juliet is a large teaching suite from NIST. OWASP Benchmark is a scored Java suite. Securibench Micro is a small set of servlet examples. Find Security Bugs contributes detector test code. Vul4J and the CVEfixes slice are closer to real project bugs, but we only keep the Java files that match the weaknesses we study.",
            "We ran the pattern checker on every file we ingested. F1 is again the combined score of false alarms and missed bugs. The best F1 here is 0.435, on 79 Find Security Bugs files. The lowest is 0.072, on 119 Securibench files. Juliet is the largest set, 20,728 files, and its F1 is 0.316.",
            "A larger set is not an easier set. The checker is a list of text patterns. It is not CodeQL, and it is not a commercial scanner that follows data through the program.",
            "We did run the live model earlier on small samples of these sets. Those scores are withdrawn. The questions still contained giveaway words, such as a method named bad or a comment that said the code was flawed. A fair test must not show the answer inside the question.",
            "The public sets also do not all contain every weakness we study. When a set has no files for a weakness, we say \"not present\". We do not rename a nearby weakness to fill the gap.",
        ],
        "example": [
            "Juliet has thousands of small teaching files. A file name often says whether it is the bad version or the good version. If we leave the word bad in the question, a model can cheat. That is why the old model scores are withdrawn.",
            "Securibench is small, 119 files, and the simple checker almost never matches the labels. F1 0.072 means the pattern list does not fit those servlets. A bigger Juliet score of 0.316 is still only a modest pattern match, not a success.",
            "Find Security Bugs is the highest bar in the picture, 0.435, and it is only 79 files. We do not average the six rows. Each set is its own result.",
        ],
        "caption": "Each bar is the simple checker's F1. Longer is better. None of these bars is a model score.",
        "headers": ["Test set", "Files", "Simple checker F1"],
        "rows": [
            ["Juliet Java", "20728", "0.316"],
            ["OWASP Benchmark", "2740", "0.395"],
            ["Securibench Micro", "119", "0.072"],
            ["Find Security Bugs", "79", "0.435"],
            ["Vul4J", "62", "0.244"],
            ["CVEfixes Java slice", "92", "0.207"],
        ],
        "table_note": "Read each row on its own. Do not average them into one project score. Do not mix them with the 0.923 score on our 24 files.",
        "outro": [
            "The pattern checker is weak on public code. That is a real result, and it is the fair baseline.",
            "The next measurement is a live-model score on Juliet after the giveaway words are hidden. We will use matched unsafe and safe files, not a single mixed file.",
        ],
        "not_claim": "Do not quote the withdrawn model scores. The simple checker is not CodeQL and not a commercial product.",
    },
    {
        "slug": "09_juliet_pairs",
        "name": "A Juliet score we can stand behind",
        "figure": fig_pair,
        "intro": [
            "Juliet is a large public Java suite. The simple checker scores 0.316 there. An earlier model score on Juliet was withdrawn because the question showed the answer.",
            "This note is the replacement. The model sees cleaned code. Comments that say \"flaw\" are blanked, and obvious names such as bad and good are renamed. The true label is used only when we score the answer, not when we ask the question.",
        ],
        "did": [
            "A Juliet pair is one unsafe file and the matching safe file. The model is right only if it calls the bad file unsafe and the good file safe. If it says \"not sure\" on either file, the pair is wrong.",
            "We sampled pairs from single-file examples, three pairs from each weakness family that Juliet actually contains, with a fixed random seed so the sample can be repeated. Some Juliet examples hide the risky line in a second file. Those were left out, because the model would never see that line.",
            "Seventeen pairs were complete. The model got 9 right. That is 0.529. Seventeen is a small sample, so we reshuffled those 17 results many times. The score usually lands between 0.294 and 0.765. The range is wide. A single 0.529 should not be treated as a precise ranking against other papers.",
            "One more pair was cut off. The free daily token limit stopped the run after 35 of 36 files. We leave that pair out of the 9-of-17 count. It is a stopped run, not a model mistake. The status of the model row is partial.",
            "If we ignore the pair rule and score each file alone, F1 on the 35 scored files is 0.789. That number is easier to inflate, because a model can mark both files unsafe and still look good on the unsafe ones. The pair score is the one we stand behind.",
        ],
        "example": [
            "Take one pair. The bad file builds a query in an unsafe way. The good file does the safe version of the same lesson. If the model says both are unsafe, it got the bad file right and the good file wrong. The pair counts as wrong.",
            "If it says the bad file is unsafe and the good file is safe, the pair counts as right. That is one of the 9.",
            "The range 0.294 to 0.765 means a different draw of 17 pairs could have looked much better or much worse. We show the range so 0.529 is not treated as exact.",
        ],
        "caption": "Both files in a pair have to be judged correctly. One right answer is not enough.",
        "headers": ["Item", "Value"],
        "rows": [
            ["Complete pairs correct", "9 of 17"],
            ["Pair accuracy", "0.529"],
            ["Likely range if the sample is repeated", "0.294 to 0.765"],
            ["Files scored before the token limit", "35 of 36"],
            ["F1 if each file is scored alone", "0.789"],
        ],
        "table_note": "Use 9 of 17, with the range, when talking about Juliet. Do not replace it with the 0.789 file-level F1, and do not mix it with the 0.923 score on our own 24 files.",
        "outro": [
            "This is the public model number we can defend today. It is modest, and the range is wide.",
            "We still have not run these same pairs with CWE search turned off. Until that run exists, we cannot say whether the book helps on Juliet. A local screen for running these tests is the practical next step, together with that missing comparison.",
        ],
        "not_claim": "We do not say that CWE search helps on Juliet. We do not say this is the best published detector.",
    },
    {
        "slug": "10_how_to_run",
        "name": "How to run the tool, and what is still open",
        "figure": fig_console,
        "intro": [
            "The tool can explain a Java method with a CWE page, and we have scores: a perfect score only on 12 practice files, a 0.923 score on 24 hard files we wrote, modest pattern-checker scores on six public sets, and a Juliet pair score of 9 out of 17.",
            "This note is about using that tool without memorizing commands, and about the tests that are still not done. It does not add a new score.",
        ],
        "did": [
            "A web page on this computer lists the test sets, the settings, and a token budget. Starting a job runs the same code as the command line. The page shows the exact command before the job starts, so the run can be repeated later.",
            "The page listens only on this machine, at 127.0.0.1. It has no accounts, because it is not open to a network. A missing model key stops a live run and says which setting is missing. It does not invent a score. An offline setting still runs, using the template instead of the live model.",
            "The six public Java sets are stored in the project, with a note of where they came from and which license applies. A fresh copy of the project can score the simple checker without downloading them again.",
            "When code is pushed, GitHub runs the unit tests. Those tests do not call Groq. They check the form, the book, the sanitizer, and the scoring math.",
            "One job runs at a time. The free model key has a daily token cap, and two big runs would spend it twice.",
        ],
        "example": [
            "Open the page on this computer. Choose the seed pipeline, the test split, and the offline setting. The page prints the command it will run. The result is four practice files, scored by the template, with no call to Groq.",
            "Choose a live Juliet run without a key and the page stops with the same message as the command line. It does not fill in a score.",
            "The unfinished Juliet file, the search-off Juliet run, the live Vul4J slices, a second model, and the human explanation labels are still open. They need a key, tokens, or a person. This note does not close them.",
        ],
        "caption": "The page is only a front door. The steps behind it are the same pipeline: mark lines, search, write, check.",
        "headers": ["Still open", "Why it matters", "Status"],
        "rows": [
            ["Finish the Juliet file cut off by the token limit", "The pair score is missing one file.", "Not finished"],
            ["Same Juliet pairs with search turned off", "We cannot yet say whether the book helps.", "Not run"],
            ["Sliced real bugs from Vul4J with the live model", "Public pattern-checker scores are not a model result.", "Not run"],
            ["A second model", "We only have one live model so far.", "Not run"],
            ["A person marking explanations", "Detection scores do not say the explanation was good.", "Labels empty"],
        ],
        "table_note": "Empty explanation labels stay empty. We will not fill them in to make the table look finished, and we have no agreement score between two people.",
        "outro": [
            "Anyone with the project can open the page, pick an offline run, and see a note. A live run still needs a Groq key and the daily token budget.",
            "The honest summary is unchanged by this screen: the model helps on our hard examples, the pattern checker is weak on public code, and the Juliet pair score is 9 of 17 with a wide range.",
        ],
        "not_claim": "A screen is not a new scientific result. We still do not say that CWE search improves Juliet detection.",
    },
]


def _shade(cell, hex_color: str) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def _set_run(run, text: str, *, size: int, bold: bool = False, color: RGBColor | None = None) -> None:
    run.text = text
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    if color is not None:
        run.font.color.rgb = color


def _add_text(doc: Document, text: str, *, size: int = 12, bold: bool = False, color: RGBColor | None = None, space_after: int = 6) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.space_before = Pt(0)
    _set_run(paragraph.add_run(), text, size=size, bold=bold, color=color)


def _build(report: dict, figure: Path) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(12)

    _add_text(doc, TITLE, size=14, bold=True, color=RGBColor(0x1B, 0x3A, 0x4B), space_after=2)
    _add_text(doc, WHO, size=11, color=RGBColor(0x5C, 0x56, 0x4C), space_after=2)
    _add_text(doc, report["name"], size=16, bold=True, color=RGBColor(0x1C, 0x19, 0x15), space_after=6)

    _add_text(doc, "Where this fits", size=13, bold=True, color=RGBColor(0x1B, 0x3A, 0x4B), space_after=4)
    for sentence in report["intro"]:
        _add_text(doc, sentence, size=12, space_after=4)

    _add_text(doc, "What we did", size=13, bold=True, color=RGBColor(0x1B, 0x3A, 0x4B), space_after=4)
    for sentence in report["did"]:
        _add_text(doc, sentence, size=12, space_after=4)

    _add_text(doc, "A concrete example", size=13, bold=True, color=RGBColor(0x1B, 0x3A, 0x4B), space_after=4)
    for sentence in report["example"]:
        _add_text(doc, sentence, size=12, space_after=4)

    _add_text(doc, "Picture", size=13, bold=True, color=RGBColor(0x1B, 0x3A, 0x4B), space_after=4)
    doc.add_picture(str(figure), width=Inches(6.4))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_text(doc, report["caption"], size=11, color=RGBColor(0x5C, 0x56, 0x4C), space_after=6)

    _add_text(doc, "The numbers", size=13, bold=True, color=RGBColor(0x1B, 0x3A, 0x4B), space_after=4)
    table = doc.add_table(rows=1, cols=len(report["headers"]))
    table.style = "Table Grid"
    for index, header in enumerate(report["headers"]):
        cell = table.rows[0].cells[index]
        cell.text = ""
        _set_run(cell.paragraphs[0].add_run(), header, size=11, bold=True, color=RGBColor(0xFF, 0xFC, 0xF7))
        _shade(cell, "1B3A4B")
    for row in report["rows"]:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = ""
            _set_run(cells[index].paragraphs[0].add_run(), value, size=11)

    _add_text(doc, report["table_note"], size=12, space_after=6)
    _add_text(doc, "Closing", size=13, bold=True, color=RGBColor(0x1B, 0x3A, 0x4B), space_after=4)
    for sentence in report["outro"]:
        _add_text(doc, sentence, size=12, space_after=4)
    _add_text(doc, "What we are not saying", size=13, bold=True, color=RGBColor(0x1B, 0x3A, 0x4B), space_after=2)
    _add_text(doc, report["not_claim"], size=12, space_after=2)
    return doc


def write_pdf(report: dict, figure: Path) -> tuple[Path, int]:
    from fpdf import FPDF

    pdf = FPDF(format="Letter")
    pdf.set_auto_page_break(auto=True, margin=16)
    pdf.set_margins(18, 16, 18)
    pdf.add_page()

    def say(text: str, size: int, bold: bool = False, color: tuple[int, int, int] = (28, 25, 21), height: float = 6) -> None:
        pdf.set_x(pdf.l_margin)
        pdf.set_text_color(*color)
        pdf.set_font("Helvetica", "B" if bold else "", size)
        pdf.multi_cell(0, height, text, new_x="LMARGIN", new_y="NEXT")

    say(TITLE, 13, bold=True, color=(27, 58, 75))
    pdf.ln(1)
    say(WHO, 11, color=(92, 86, 76))
    pdf.ln(1)
    say(report["name"], 16, bold=True, height=8)
    pdf.ln(1)
    say("Where this fits", 13, bold=True, color=(27, 58, 75), height=7)
    for sentence in report["intro"]:
        say(sentence, 12)
        pdf.ln(1)
    say("What we did", 13, bold=True, color=(27, 58, 75), height=7)
    for sentence in report["did"]:
        say(sentence, 12)
        pdf.ln(1)
    say("A concrete example", 13, bold=True, color=(27, 58, 75), height=7)
    for sentence in report["example"]:
        say(sentence, 12)
        pdf.ln(1)
    pdf.ln(1)
    say("Picture", 13, bold=True, color=(27, 58, 75), height=7)
    pdf.image(str(figure), w=160)
    pdf.ln(2)
    say(report["caption"], 11, color=(92, 86, 76))
    pdf.ln(2)
    say("The numbers", 13, bold=True, color=(27, 58, 75), height=7)
    pdf.set_font("Helvetica", size=11)
    pdf.set_text_color(28, 25, 21)
    with pdf.table(width=170) as table:
        heading = table.row()
        for header in report["headers"]:
            heading.cell(header)
        for values in report["rows"]:
            row = table.row()
            for value in values:
                row.cell(value)
    pdf.ln(2)
    say(report["table_note"], 12)
    pdf.ln(2)
    say("Closing", 13, bold=True, color=(27, 58, 75), height=7)
    for sentence in report["outro"]:
        say(sentence, 12)
        pdf.ln(1)
    say("What we are not saying", 13, bold=True, color=(27, 58, 75), height=7)
    say(report["not_claim"], 12)
    dest = OUT / f"{report['slug']}_RichaVerma_25MCSS02.pdf"
    pages = len(pdf.pages)
    pdf.output(dest)
    return dest, pages


def write_readme() -> None:
    lines = [
        "# Progress reports for the mentor",
        "",
        "Ten separate notes. Each one can be read on its own. The opening and the closing connect it to the work around it.",
        "The Word file is the one to edit. The PDF has the same words and the same picture, for sending.",
        "",
        "Student: Richa Verma (25MCSS02). Advisor: Dr. Akshay Pandey.",
        "",
    ]
    for report in REPORTS:
        lines.append(f"- {report['name']}. {report['intro'][0]}")
    lines.append("")
    lines.append("Regenerate with `uv run python scripts/write_progress_reports.py`.")
    lines.append("")
    (OUT / "README.md").write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    for report in REPORTS:
        figure = report["figure"]()
        doc = _build(report, figure)
        dest = OUT / f"{report['slug']}_RichaVerma_25MCSS02.docx"
        doc.save(dest)
        pdf, pages = write_pdf(report, figure)
        print(dest)
        print(pdf, "pages", pages)
    write_readme()
    print(OUT / "README.md")


if __name__ == "__main__":
    main()
