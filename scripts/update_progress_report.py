"""Update the thesis progress report DOCX in place (template headings kept)."""

from __future__ import annotations

from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "RAG_MultiAgent_Vulnerability_Thesis_Progress_Report.docx"
OFFICIAL = ROOT / "Thesis Progress Report_RichaVerma_25MCSS02.docx"


def set_para(doc: Document, index: int, text: str) -> None:
    para = doc.paragraphs[index]
    if para.runs:
        para.runs[0].text = text
        for run in para.runs[1:]:
            run.text = ""
    else:
        para.add_run(text)


def set_cell(cell, text: str) -> None:
    if cell.paragraphs and cell.paragraphs[0].runs:
        cell.paragraphs[0].runs[0].text = text
        for run in cell.paragraphs[0].runs[1:]:
            run.text = ""
        for extra in cell.paragraphs[1:]:
            extra.clear()
    elif cell.paragraphs:
        cell.paragraphs[0].text = text
    else:
        cell.text = text


def main() -> None:
    doc = Document(str(SRC))

    set_para(
        doc,
        1,
        "RICHA VERMA – 25MCSS02  ·  Advisor: Dr. Akshay Pandey",
    )

    set_para(
        doc,
        28,
        "Assignments 1–4 and the wired seed pipeline are implemented (PRs #1–#10), "
        "then MiniLM embeddings and a Groq LLM reasoner (embeddings-llm). Research "
        "evaluation is complete on an authored expanded Java corpus (36 units: original "
        "12 plus 24 held-out FP/FN traps) — not a public benchmark (not Juliet, OWASP "
        "Benchmark, or Big-Vul). Live Groq openai/gpt-oss-20b was used in trials. "
        "Without GROQ_API_KEY the default path is offline. No OpenAI account is required.",
    )

    set_para(
        doc,
        30,
        "The repository now uses layered subpackages under src/cwe_vuln/ (models, dataset, "
        "knowledge, sast, retrieval, schema, reasoner, validator, orchestrator, framework, cli). "
        "Shared DTOs live in models/. Tests mirror those packages. data/ remains the corpus; "
        "Java seeds are not inside the Python package.",
    )

    set_para(
        doc,
        39,
        "The implemented architecture combines regex SAST evidence, hybrid CWE retrieval "
        "(MiniLM cosine with TF-IDF fallback, SAST ids, CWE relationships, RRF), a cost-aware "
        "orchestrator, a template reasoner, an LLMReasoner behind the same port, and a "
        "validator. Without GROQ_API_KEY (or CWE_VULN_LLM_API_KEY override) the LLM is not "
        "constructed. No OpenAI account is required; OPENAI_API_KEY is unused.",
    )

    set_para(
        doc,
        42,
        "The SAST layer runs seed-oriented regex rules and emits Evidence objects (file, lines, "
        "snippet, CWE id hint). It does not embed CWE encyclopedia text; names and mitigations "
        "come from the knowledge store. A “safe” regex result is not treated as a hard proof of "
        "absence — the pipeline still retrieves CWE context for explanation.",
    )

    set_para(
        doc,
        43,
        "The hybrid retrieval layer ranks curated CWE entries using MiniLM cosine when the local "
        "all-MiniLM-L6-v2 model is available, otherwise TF-IDF, plus SAST-matched CWE ids and "
        "one-hop CWE relationships, fused with Reciprocal Rank Fusion. If MiniLM cannot load, "
        "results record embedder=tfidf_fallback. Retrieval confidence is not a vulnerability decision.",
    )

    set_para(
        doc,
        44,
        "The cost-aware orchestrator runs SAST first. If no GROQ_API_KEY or CWE_VULN_LLM_API_KEY "
        "is set, the LLM is skipped and TemplateReasoner runs. With a key, the "
        "path is sast_then_llm or hybrid_retrieve_then_llm unless skip_llm_when_sast_hits is set. "
        "Default model is Groq openai/gpt-oss-20b at https://api.groq.com/openai/v1. "
        "No OpenAI account is required; OPENAI_API_KEY is unused.",
    )

    set_para(
        doc,
        45,
        "The validator checks schema validity, that the cited CWE exists in the knowledge store, "
        "that supporting source lines exist in the Java unit, and that the decision is consistent "
        "with SAST evidence. Framework metrics compare decisions to seed labels.",
    )

    set_para(
        doc,
        55,
        "The seed pipeline is implemented and runnable: uv sync && uv run pytest && "
        "uv run cwe-vuln-pipeline. Assignment 1 detection on 12 units P=R=F1=1.000 "
        "(seed-only). Research evaluation on 24 held-out authored traps: SAST/template "
        "P=R=F1=0.000 (12 FP / 12 FN); live Groq then_llm P=0.857 R=1.000 F1=0.923; "
        "skip_llm P=0.500 R=1.000 F1=0.667. Retrieval on 48 authored queries: hybrid "
        "R@1=0.854 MRR=0.917 vs MiniLM 0.812 / 0.894 vs TF-IDF 0.646 / 0.794. "
        "Public-benchmark evaluation is not claimed.",
    )

    # Section 12 next steps: remaining honest gaps
    set_para(doc, 57, "Keep the original 12-unit Java seed as the assignment split; research scores use the 24-unit held-out authored traps.")
    set_para(doc, 58, "Optional later: expand the curated CWE store without scraping an unofficial full dump.")
    set_para(doc, 59, "MiniLM embeddings are implemented with TF-IDF fallback if the local model is missing.")
    set_para(doc, 60, "Keep SAST as evidence extraction (CWE hints); knowledge layer owns descriptions.")
    set_para(doc, 61, "Keep SAST-first LLM routing in the orchestrator (GROQ_API_KEY).")
    set_para(doc, 62, "Live LLM reasoner is implemented; research trials used Groq openai/gpt-oss-20b.")
    set_para(doc, 63, "Validator already checks schema, CWE id, cited lines, and decision vs evidence — it currently fails when the LLM correctly overrides SAST.")
    set_para(doc, 64, "Optional later: richer confidence fusion beyond template confidence fields.")
    set_para(doc, 65, "Public-benchmark experiments remain out of scope until a labeled external dataset is adopted.")

    set_para(
        doc,
        70,
        "The thesis has progressed from study of Explainable AI, RAG and agentic systems to a "
        "runnable seed-only pipeline for explainable CWE-oriented vulnerability detection. The "
        "implemented path combines a curated CWE knowledge store, MiniLM/TF-IDF hybrid retrieval, "
        "regex SAST evidence, a template reasoner, an LLM path gated on GROQ_API_KEY or "
        "CWE_VULN_LLM_API_KEY, validation, and cost-aware orchestrator routing.",
    )

    set_para(
        doc,
        71,
        "Current work establishes the problem, layered architecture, MiniLM embeddings, "
        "an LLM reasoner, and a completed research evaluation on an authored expanded "
        "corpus (with limitations: small N, traps written against our regex). "
        "Public-benchmark evaluation remains future work and is not reported as done.",
    )

    # Table 4 — area / status / progress
    t4 = doc.tables[4]
    rows = [
        ("Area", "Status", "Current progress"),
        ("XAI and explainability", "Studied", "Covered explainability, SHAP, counterfactuals and faithfulness-related concepts."),
        ("RAG", "Studied", "Understood retrieval, grounding and the role of external knowledge."),
        ("Agentic AI", "Studied", "Defined specialized agents and orchestration."),
        ("CWE security knowledge", "Implemented", "Curated store + query API (PR #2). Teaching subset, not a MITRE dump."),
        ("Vulnerability detection problem", "Defined", "Focused on precision, false positives, false negatives and explainability."),
        ("Static analysis role", "Implemented", "Regex SAST + Evidence spans (PRs #1, #5). Seed-only baseline P=R=F1=1.000."),
        ("Hybrid retrieval", "Implemented", "MiniLM cosine + SAST + CWE relationships, RRF. TF-IDF fallback if MiniLM missing."),
        ("Cost-aware routing", "Implemented", "Orchestrator SAST-first; Groq LLM only if GROQ_API_KEY is set (CWE_VULN_LLM_API_KEY override)."),
        ("Multi-agent workflow", "Implemented", "Template + LLM reasoner, validator, framework CLI."),
        ("Implementation / evaluation", "Research eval done", "Authored 36-unit corpus. Live Groq trials recorded. Not a public benchmark."),
    ]
    for i, (a, b, c) in enumerate(rows):
        set_cell(t4.rows[i].cells[0], a)
        set_cell(t4.rows[i].cells[1], b)
        set_cell(t4.rows[i].cells[2], c)

    tree = (
        "src/cwe_vuln/\n"
        "  config.py\n"
        "  models/          Evidence, RankedHit, ReasoningResult, ValidationReport, metrics\n"
        "  dataset/         SeedUnit, load_seed 8/4, load_research_corpus 36\n"
        "  knowledge/       CWEKnowledgeBase\n"
        "  sast/            regex detect + extract_evidence\n"
        "  retrieval/       MiniLM Embedder + TF-IDF fallback + SAST + RRF\n"
        "  schema/          JSON Schema validate helpers\n"
        "  reasoner/        TemplateReasoner + LLMReasoner\n"
        "  validator/       ResultValidator\n"
        "  orchestrator/    Pipeline, LLM routing if API key present\n"
        "  framework/       cwe-vuln-pipeline + cwe-vuln-eval\n"
        "  cli/             baseline, kb, retrieve, schema, evidence\n"
        "data/  docs/  results/  schemas/  tests/ (mirrors packages)"
    )
    set_cell(doc.tables[5].rows[0].cells[0], tree)

    t8 = doc.tables[8]
    rows8 = [
        ("Component", "Status", "Next step"),
        ("Research direction", "Defined", "Continue literature study around vulnerability detection, RAG and multi-agent systems."),
        ("Problem formulation", "Completed", "Keep evaluation criteria seed-only until an external dataset is adopted."),
        ("CWE knowledge layer", "Implemented", "Optional: enlarge the curated subset; do not scrape a messy dump."),
        ("Static analysis", "Implemented", "Regex evidence extraction on the 12-unit Java seed."),
        ("Hybrid retrieval", "Implemented", "MiniLM with TF-IDF fallback; seed-only metrics in results/."),
        ("Cost-aware orchestrator", "Implemented", "SAST-first; LLM path when GROQ_API_KEY or CWE_VULN_LLM_API_KEY is set."),
        ("Reasoning agent", "Implemented", "Template offline; LLMReasoner (default openai/gpt-oss-20b)."),
        ("Validator agent", "Implemented", "Schema, CWE id, cited lines, decision vs evidence."),
        ("Layered packages", "Implemented", "models, dataset, knowledge, sast, retrieval, schema, reasoner, validator, orchestrator, framework, cli."),
        ("Reporting", "Implemented", "Assignment 4 JSON Schema + framework results JSON/MD."),
        ("Evaluation", "Authored corpus", "36 units / 24 held-out traps. No Juliet / OWASP Benchmark / Big-Vul numbers."),
    ]
    for i, (a, b, c) in enumerate(rows8):
        set_cell(t8.rows[i].cells[0], a)
        set_cell(t8.rows[i].cells[1], b)
        set_cell(t8.rows[i].cells[2], c)

    set_cell(
        doc.tables[9].rows[0].cells[0],
        "Progress summary: XAI → RAG → CWE knowledge → agentic AI → Java seed + SAST baseline "
        "→ hybrid retrieval → schema → evidence → template reasoner → validator → "
        "orchestrator → framework CLI → layered packages → MiniLM embeddings + LLM "
        "reasoner → research evaluation on an authored expanded corpus. Live Groq used "
        "in trials. Public-benchmark evaluation remains later.",
    )

    set_cell(
        doc.tables[0].rows[0].cells[0],
        "Research focus: RAG-augmented multi-agent CWE vulnerability detection. Runnable path is "
        "seed-only (regex SAST, MiniLM/TF-IDF hybrid retrieval, template reasoner, optional "
        "LLM). Cost-aware routing lives in the orchestrator.",
    )

    doc.save(str(SRC))
    doc.save(str(OFFICIAL))
    print(f"wrote {SRC}")
    print(f"wrote {OFFICIAL}")


if __name__ == "__main__":
    main()
